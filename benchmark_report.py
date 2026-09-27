"""Measured benchmark data -> publication-style plots and a numeric Excel workbook.

This module never imports Torch or runs inference. Failed/partial runs remain in
the table but become gaps in plots, not zero-valued or interpolated results.
"""

import csv
import json
import math
from pathlib import Path


FIELDS = (
    "prune_rate", "gpu", "gpu_name", "status", "complete", "exit_code",
    "selected_samples", "evaluated_samples", "labeled_samples", "correct", "incorrect", "unparsed",
    "accuracy", "accuracy_percent", "total_seconds", "worker_wall_seconds",
    "generation_seconds_sum", "mean_generation_ms", "generation_images_per_second",
    "peak_allocated_gib", "peak_reserved_gib", "peak_allocated_bytes", "peak_reserved_bytes",
    "error", "memory_error", "result_read_error", "output_dir", "log",
)

DEFINITIONS = {
    "prune_rate": "Requested pruning percentage, not measured token removal percentage.",
    "gpu": "Physical GPU ID used by this job; differing GPUs/load can confound timing comparisons.",
    "gpu_name": "Device name reported by PyTorch.",
    "status": "Only success + complete + matching sample counts are plotted; failures are retained in the table.",
    "complete": "Whether evaluation finished. Failure status takes priority even if this is true.",
    "exit_code": "Worker process exit code; 0 indicates normal exit.",
    "selected_samples": "First min(100, input record count) samples, in source-file order, shared by every rate.",
    "evaluated_samples": "Number of samples processed; incomplete jobs must not be compared as complete runs.",
    "labeled_samples": "Number of evaluated samples with gt; denominator used for accuracy.",
    "correct": "Correct labeled predictions (A=defect, B=no defect).",
    "incorrect": "Labeled samples minus correct; includes unparsed answers.",
    "unparsed": "Labeled predictions without a leading A/B option.",
    "accuracy": "Raw fraction, e.g. 0.9 = 90%; null when there are no labels.",
    "accuracy_percent": "100 * accuracy for successful complete runs; percentage, not AUROC or F1.",
    "total_seconds": "Process launch to exit: includes model loading, preprocessing, inference and JSON/log writes; excludes queue wait and final chart/Excel export. Polling resolution 0.2 s.",
    "worker_wall_seconds": "Worker wall time through evaluation, before its final resource report is written.",
    "generation_seconds_sum": "Sum of synchronized per-image generate() seconds; includes pruning Q/K scoring and the first cold call. Excludes model loading, preprocessing and result writes; no warmup exclusion.",
    "mean_generation_ms": "1000 * generation_seconds_sum / evaluated_samples; per-image response time, NOT per-token latency or TTFT.",
    "generation_images_per_second": "evaluated_samples / generation_seconds_sum; generation-only images/s, not end-to-end throughput.",
    "peak_allocated_gib": "PyTorch peak tensor allocations across model load + selected samples, in GiB (1024^3 bytes).",
    "peak_reserved_gib": "PyTorch peak reserved pool including cache across model load + selected samples; not whole-board nvidia-smi usage.",
    "peak_allocated_bytes": "Same allocated peak in bytes; excludes allocations outside PyTorch.",
    "peak_reserved_bytes": "Same reserved peak in bytes; excludes CUDA contexts and allocations outside PyTorch.",
    "error": "Job failure reason, if any.",
    "memory_error": "Resource measurement failure, if any; missing values are not zero.",
    "result_read_error": "Result-file read error, if any.",
    "output_dir": "Job output directory containing predictions and accuracy metrics.",
    "log": "Worker stdout/stderr log path.",
}


def check_report_dependencies():
    try:
        import matplotlib
        matplotlib.use("Agg")  # Headless Linux servers; no GUI/display required.
        import matplotlib.pyplot  # noqa: F401
        import openpyxl  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError as error:
        raise RuntimeError(
            "Chart/Excel dependencies are missing. Run: python -m pip install -r requirements-report.txt"
        ) from error


def number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return value
    return None


def completed(row):
    return (row.get("status") == "success" and row.get("complete") is True
            and row.get("exit_code", 0) == 0
            and (row.get("selected_samples") is None or
                 row.get("evaluated_samples") == row["selected_samples"]))


def metric_rows(summary):
    jobs = summary.get("jobs")
    if not isinstance(jobs, list) or not jobs:
        raise ValueError("benchmark.json must contain a nonempty jobs list")
    rates = [job.get("prune_rate") for job in jobs]
    if any(type(rate) is not int or not 0 <= rate <= 90 for rate in rates) or len(set(rates)) != len(rates):
        raise ValueError("Benchmark rates must be unique integers in [0, 90]")
    rows = []
    for job in sorted(jobs, key=lambda job: job["prune_rate"]):
        row = dict(job, accuracy_percent=None, mean_generation_ms=None, generation_images_per_second=None)
        if completed(row):
            accuracy = number(row.get("accuracy"))
            if accuracy is not None:
                row["accuracy_percent"] = accuracy * 100
            seconds = number(row.get("generation_seconds_sum"))
            count = number(row.get("evaluated_samples"))
            if seconds is not None and seconds > 0 and count is not None and count > 0:
                row["mean_generation_ms"] = seconds * 1000 / count
                row["generation_images_per_second"] = count / seconds
        rows.append(row)
    return rows


def curve_values(rows, key):
    # Keep failed rates on the x axis but break the line at them.
    return [value if completed(row) and (value := number(row.get(key))) is not None
            else float("nan") for row in rows]


def plot_metrics(rows, summary, output):
    import matplotlib.pyplot as plt
    from matplotlib.ticker import MaxNLocator, ScalarFormatter

    blue, red = "#0876B9", "#E5262D"
    panels = (
        ("Image-level Accuracy", "Accuracy (%)", (("accuracy_percent", "Accuracy", blue, "o", "-"),)),
        ("Elapsed Time", "Time (s)", (("total_seconds", "Total", blue, "o", "-"),
                                     ("generation_seconds_sum", "Generation", red, "s", "--"))),
        ("Peak GPU Memory", "Memory (GiB)", (("peak_allocated_gib", "Allocated", blue, "o", "-"),
                                            ("peak_reserved_gib", "Reserved", red, "s", "--"))),
        ("Mean Generation Time", "Time per image (ms)", (("mean_generation_ms", "Generation / image", blue, "o", "-"),)),
    )
    style = {"font.family": "serif", "font.serif": ["Times New Roman", "DejaVu Serif"],
             "font.size": 11, "axes.titlesize": 13, "axes.titleweight": "bold",
             "axes.labelsize": 12, "axes.edgecolor": "black", "axes.linewidth": 0.9,
             "figure.facecolor": "white", "axes.facecolor": "white", "pdf.fonttype": 42}
    with plt.rc_context(style):
        fig, axes = plt.subplots(1, 4, figsize=(18.4, 4.9))
        try:
            fig.subplots_adjust(left=0.055, right=0.987, bottom=0.25, top=0.78, wspace=0.32)
            for ax, (title, ylabel, series) in zip(axes, panels):
                has_data = False
                for key, label, color, marker, linestyle in series:
                    values = curve_values(rows, key)
                    has_data |= any(math.isfinite(value) for value in values)
                    ax.plot([row["prune_rate"] for row in rows], values, color=color,
                            marker=marker, linestyle=linestyle, linewidth=1.8, markersize=4.7, label=label)
                ax.set(title=title, xlabel="Pruning rate (%)", ylabel=ylabel, xlim=(-3, 93))
                ax.set_xticks(list(range(0, 100, 10)))
                ax.tick_params(labelsize=9.5, direction="out", length=3)
                ax.grid(True, color="#B8B8B8", linestyle=":", linewidth=0.6, alpha=0.85)
                ax.set_axisbelow(True)
                ax.yaxis.set_major_locator(MaxNLocator(nbins=6))
                formatter = ScalarFormatter(useOffset=False)
                formatter.set_scientific(False)
                ax.yaxis.set_major_formatter(formatter)
                ax.legend(loc="lower center", bbox_to_anchor=(0.5, 1.15), frameon=False,
                          ncol=len(series), fontsize=10, handlelength=2.8, columnspacing=1.1)
                if not has_data:
                    ax.text(0.5, 0.5, "No complete data", transform=ax.transAxes, ha="center", color="#666666")
            settings = summary.get("settings", {})
            count = settings.get("selected_samples", "selected")
            gpus = sorted({row["gpu"] for row in rows if row.get("gpu") is not None})
            failed = sum(not completed(row) for row in rows)
            note = (f"First {count} samples | GPUs: {', '.join(map(str, gpus))} | "
                    "Greedy decoding | No sample visualizations | No warmup exclusion")
            if len(gpus) > 1:
                note += "\nDifferent GPUs and concurrent load may affect timing comparisons."
            if failed:
                note += f"\n{failed} failed/incomplete rate(s): gaps in curves; raw values retained in Excel."
            fig.text(0.5, 0.035, note, ha="center", va="bottom", fontsize=9, color="#444444")
            fig.savefig(output / "benchmark_curves.png", dpi=300, facecolor="white")
            fig.savefig(output / "benchmark_curves.pdf", facecolor="white")
        finally:
            plt.close(fig)


def excel_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_workbook(rows, summary, output):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(list(FIELDS))
    for row in rows:
        sheet.append([excel_value(row.get(key)) for key in FIELDS])
    sheet.freeze_panes = "C2"
    table = Table(displayName="BenchmarkMetrics", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    for index, field in enumerate(FIELDS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(32, max(16, len(field) + 2))
        for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            if field == "accuracy":
                cells[0].number_format = "0.00%"
            elif field in {"accuracy_percent", "total_seconds", "worker_wall_seconds", "generation_seconds_sum",
                           "mean_generation_ms", "generation_images_per_second", "peak_allocated_gib", "peak_reserved_gib"}:
                cells[0].number_format = "0.0000"
    settings = workbook.create_sheet("Settings")
    settings.append(["Setting", "Value"])
    for name in ("experiment_wall_seconds", "time_scope", "generation_time_scope", "memory_scope"):
        settings.append([name, excel_value(summary.get(name))])
    for name, value in summary.get("settings", {}).items():
        settings.append([name, excel_value(value)])
    definitions = workbook.create_sheet("Definitions")
    definitions.append(["Column", "Meaning / unit"])
    for name in FIELDS:
        definitions.append([name, DEFINITIONS[name]])
    for tab in (settings, definitions):
        tab.column_dimensions["A"].width = 32
        tab.column_dimensions["B"].width = 115
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cells in tab.iter_rows(min_row=2):
            cells[1].alignment = Alignment(wrap_text=True, vertical="top")
            tab.row_dimensions[cells[1].row].height = 45
    for tab in (sheet, settings, definitions):
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="126DA5")
        # Source strings (paths/errors) are data, never spreadsheet formulas.
        for cells in tab.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    figure = workbook.create_sheet("Figure")
    figure["A1"] = "Measured metrics versus pruning rate; see Definitions for units and timing scope."
    picture = Image(str(output / "benchmark_curves.png"))
    picture.width, picture.height = 1472, 392
    figure.add_image(picture, "A3")
    workbook.save(output / "benchmark.xlsx")
    workbook.close()


def generate_reports(summary_path):
    """Regenerate derived exports only; leave source JSON, logs and predictions untouched."""
    summary_path = Path(summary_path)
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    rows = metric_rows(summary)
    check_report_dependencies()
    output = summary_path.parent
    with (output / "benchmark.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    plot_metrics(rows, summary, output)
    write_workbook(rows, summary, output)
    return rows
