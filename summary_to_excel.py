"""Export an existing summary.csv to Excel without running model inference.

Edit SUMMARY_CSV below and run: python summary_to_excel.py
Or run: python summary_to_excel.py /path/to/summary.csv
Only Python and openpyxl are required. Source CSV and existing exports are kept.
"""

# ===== 在这里填写 summary.csv 的路径，然后直接运行本文件 =====
# 绝对路径可直接填写；这里的相对路径以本脚本所在目录为起点。
SUMMARY_CSV = r"output/vico_exclude_time_01/summary.csv"

import argparse
import csv
import io
import math
from pathlib import Path


PERCENT_COLUMNS = {"accuracy", "precision", "recall", "tnr"}
COUNT_COLUMNS = {
    "expected_samples", "evaluated_samples", "labeled_samples", "correct", "incorrect",
    "unparsed", "parsed_samples", "tp", "fp", "tn", "fn",
}


def configured_input():
    if not isinstance(SUMMARY_CSV, str) or not SUMMARY_CSV.strip():
        raise ValueError("Set SUMMARY_CSV at the top of this script to your summary.csv path")
    source = Path(SUMMARY_CSV).expanduser()
    if not source.is_absolute():
        source = Path(__file__).resolve().parent / source
    return source.resolve()


def _cell_value(column, raw):
    if raw == "":
        return None
    if column == "complete":
        if raw.strip().lower() not in {"true", "false"}:
            raise ValueError("complete must be True, False or empty")
        return raw.strip().lower() == "true"
    if column in COUNT_COLUMNS or column == "prune_rate":
        value = int(raw)
        if value < 0 or (column == "prune_rate" and value > 99):
            raise ValueError(f"Invalid nonnegative count/pruning percentage: {column}={raw}")
        return value
    if column in PERCENT_COLUMNS or column.endswith("_seconds"):
        value = float(raw)
        if not math.isfinite(value) or value < 0 or (column in PERCENT_COLUMNS and value > 1):
            raise ValueError(f"{column} must be finite and nonnegative; metrics must be fractions in [0, 1]")
        return value
    # Preserve future/unknown fields literally, including leading-zero IDs.
    return raw


def _read_csv(source):
    with source.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        columns = next(reader, None)
        if (not columns or any(not name.strip() for name in columns)
                or len(set(columns)) != len(columns)):
            raise ValueError("CSV must have nonempty, unique column headers")
        if len(columns) > 16384:
            raise ValueError("CSV exceeds Excel's column limit")
        rows = []
        for record in reader:
            if not record:  # Ignore empty physical lines, not empty metric cells.
                continue
            if len(record) != len(columns):
                raise ValueError(f"CSV row ending at line {reader.line_num} has a different column count")
            try:
                rows.append([_cell_value(column, raw) for column, raw in zip(columns, record)])
            except ValueError as error:
                raise ValueError(f"Invalid CSV row ending at line {reader.line_num}: {error}") from error
            if len(rows) > 1048575:
                raise ValueError("CSV exceeds Excel's row limit (including the header)")
    if not rows:
        raise ValueError("CSV contains no data rows")
    return columns, rows


def _workbook_bytes(source, columns, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError('Install the export dependency: python -m pip install "openpyxl>=3.1,<4"') from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "summary"
    for row_number, values in enumerate([columns] + rows, start=1):
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column_number, value=value)
            if isinstance(value, str):
                # Treat CSV strings as data, never execute them as Excel formulas.
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=11)
            if row_number == 1:
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24476A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            elif columns[column_number - 1] in PERCENT_COLUMNS:
                cell.number_format = "0.00%"  # Store 0.9, display 90.00%; no multiplication.
            elif columns[column_number - 1] == "prune_rate":
                cell.number_format = '0"%"'  # Source rates already use 0..99 percentage units.
            elif columns[column_number - 1] in COUNT_COLUMNS:
                cell.number_format = "0"
            elif columns[column_number - 1].endswith("_seconds"):
                cell.number_format = "0.000000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for index, name in enumerate(columns, start=1):
        longest = max(len(name), *(len(str(row[index - 1])) if row[index - 1] is not None else 0 for row in rows))
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(13, longest + 2))
    sheet["A1"].comment = Comment(
        f"Source: {source}\nOriginal columns and row order retained; no results recalculated.\n"
        "ACC/PRE/Recall/TNR fractions are displayed as percentages; missing values remain blank.\n"
        "Unknown columns are preserved as text. Only timing columns present in the CSV are exported.",
        "summary_to_excel",
    )
    workbook.properties.title = "Pruning experiment summary"
    workbook.properties.description = f"Direct export of {source}; no model inference or metric recalculation."
    data = io.BytesIO()
    try:
        workbook.save(data)
        return data.getvalue()
    finally:
        workbook.close()


def convert_summary(input_path=None, output_path=None):
    source = configured_input() if input_path is None else Path(input_path).expanduser().resolve()
    if source.is_dir():
        source = source / "summary.csv"
    if not source.is_file():
        raise FileNotFoundError(f"summary.csv not found: {source}")
    if source.suffix.lower() != ".csv":
        raise ValueError("Input must be a CSV file")
    destination = (source.with_suffix(".xlsx") if output_path is None
                   else Path(output_path).expanduser().resolve())
    if destination == source or destination.suffix.lower() != ".xlsx":
        raise ValueError("Output must be a separate .xlsx file")
    if output_path is not None and destination.exists():
        raise FileExistsError(f"Output already exists; choose a new file: {destination}")
    columns, rows = _read_csv(source)
    data = _workbook_bytes(source, columns, rows)
    destination.parent.mkdir(parents=True, exist_ok=True)
    candidate, suffix = destination, 2
    while True:
        try:
            with candidate.open("xb") as stream:
                stream.write(data)
            return candidate
        except FileExistsError:
            if output_path is not None:
                raise
            candidate = destination.with_name(f"{destination.stem}_{suffix}.xlsx")
            suffix += 1


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("summary_csv", nargs="?", help="CSV path; overrides SUMMARY_CSV configured above")
    parser.add_argument("--output", type=Path, help="Optional new .xlsx path (never overwrites)")
    args = parser.parse_args(argv)
    try:
        destination = convert_summary(args.summary_csv, args.output)
    except (OSError, ValueError, RuntimeError, csv.Error) as error:
        parser.exit(1, f"Error: {error}\n")
    print(f"Saved Excel: {destination}")


if __name__ == "__main__":
    main()
