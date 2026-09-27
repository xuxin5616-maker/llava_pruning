import csv
import importlib.util
import json
import math
import tempfile
import unittest
from pathlib import Path

from benchmark_report import FIELDS, curve_values, generate_reports, metric_rows


def synthetic_summary():
    """Synthetic test fixture only; never represents a measured GPU experiment."""
    return {
        "settings": {"selected_samples": 100, "gpus": [4, 5], "do_sample": False,
                     "seed": 42, "model_path": "SYNTHETIC TEST FIXTURE"},
        "jobs": [{"prune_rate": rate, "gpu": 4 + (rate // 10) % 2,
                  "gpu_name": "SYNTHETIC TEST GPU", "status": "success", "complete": True,
                  "exit_code": 0, "selected_samples": 100, "evaluated_samples": 100,
                  "accuracy": 0.9 - rate / 1000, "generation_seconds_sum": 100 + rate,
                  "total_seconds": 130 + rate, "peak_allocated_gib": 19.664,
                  "peak_reserved_gib": 20.936} for rate in range(90, -1, -10)],
    }


class ReportDataTests(unittest.TestCase):
    def test_sort_units_and_derived_metrics(self):
        rows = metric_rows(synthetic_summary())
        self.assertEqual([row["prune_rate"] for row in rows], list(range(0, 100, 10)))
        self.assertEqual(rows[0]["accuracy"], 0.9)
        self.assertEqual(rows[0]["accuracy_percent"], 90)
        self.assertEqual(rows[0]["mean_generation_ms"], 1000)
        self.assertEqual(rows[0]["generation_images_per_second"], 1)

    def test_partial_failed_and_missing_measurements_are_gaps_not_zero(self):
        summary = synthetic_summary()
        summary["jobs"][0].update(status="failed", generation_seconds_sum=2)
        summary["jobs"][1].update(complete=False)
        summary["jobs"][2].update(evaluated_samples=99)
        summary["jobs"][3].update(accuracy=None, peak_allocated_gib=None, generation_seconds_sum=None)
        rows = metric_rows(summary)
        for row in rows[-3:]:
            self.assertIsNone(row["accuracy_percent"])
            self.assertIsNone(row["mean_generation_ms"])
            self.assertTrue(math.isnan(curve_values([row], "total_seconds")[0]))
        self.assertTrue(math.isnan(curve_values([rows[-4]], "accuracy_percent")[0]))
        self.assertIsNone(rows[-4]["mean_generation_ms"])

    def test_invalid_summary_does_not_export(self):
        for jobs in ([], [{"prune_rate": 0}, {"prune_rate": 0}], [{"prune_rate": "10"}]):
            with self.subTest(jobs=jobs), self.assertRaises(ValueError):
                metric_rows({"jobs": jobs})


@unittest.skipUnless(all(importlib.util.find_spec(name) for name in ("matplotlib", "openpyxl", "PIL")),
                     "requires report dependencies")
class ReportExportTests(unittest.TestCase):
    def test_plot_excel_and_csv_contain_measured_values_and_preserve_raw_json(self):
        from openpyxl import load_workbook
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            summary = synthetic_summary()
            summary["jobs"][0].update(status="failed", error="=NOT_A_FORMULA()")
            source = root / "benchmark.json"
            raw = json.dumps(summary)
            source.write_text(raw, encoding="utf-8")
            generate_reports(source)
            self.assertEqual(source.read_text(encoding="utf-8"), raw)
            for name in ("benchmark_curves.png", "benchmark_curves.pdf", "benchmark.xlsx", "benchmark.csv"):
                self.assertGreater((root / name).stat().st_size, 100)
            with Image.open(root / "benchmark_curves.png") as figure:
                self.assertGreater(figure.width, 5000)
                self.assertGreater(figure.height, 1000)
            workbook = load_workbook(root / "benchmark.xlsx")
            self.assertEqual(workbook.sheetnames, ["Metrics", "Settings", "Definitions", "Figure"])
            sheet = workbook["Metrics"]
            self.assertEqual(sheet.max_row, 11)
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers, list(FIELDS))
            value = lambda row, key: sheet.cell(row, headers.index(key) + 1)
            self.assertEqual(value(2, "prune_rate").value, 0)
            self.assertEqual(value(2, "accuracy").value, 0.9)
            self.assertEqual(value(2, "accuracy").number_format, "0.00%")
            self.assertEqual(value(2, "mean_generation_ms").value, 1000)
            self.assertEqual(value(2, "peak_allocated_gib").value, 19.664)
            self.assertIsNone(value(11, "accuracy_percent").value)
            self.assertEqual(value(11, "error").data_type, "s")
            self.assertEqual(value(11, "error").value, "=NOT_A_FORMULA()")
            workbook.close()
            with (root / "benchmark.csv").open(encoding="utf-8", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(len(rows), 10)
            self.assertEqual(rows[-1]["accuracy_percent"], "")

    def test_all_failed_still_exports_without_inventing_measurements(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            summary = synthetic_summary()
            for row in summary["jobs"]:
                row.update(status="failed", complete=False)
            source = root / "benchmark.json"
            source.write_text(json.dumps(summary), encoding="utf-8")
            rows = generate_reports(source)
            self.assertTrue(all(row["accuracy_percent"] is None for row in rows))
            self.assertTrue((root / "benchmark.xlsx").is_file())


if __name__ == "__main__":
    unittest.main()
