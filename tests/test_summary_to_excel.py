import contextlib
import csv
import importlib.util
import io
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import summary_to_excel as converter


HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None


class SummaryConfigurationTests(unittest.TestCase):
    def test_configured_absolute_and_relative_input(self):
        script_directory = Path(converter.__file__).resolve().parent
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "summary.csv"
            with patch.object(converter, "SUMMARY_CSV", str(source)):
                self.assertEqual(converter.configured_input(), source.resolve())
        with patch.object(converter, "SUMMARY_CSV", "output/custom/summary.csv"):
            self.assertEqual(converter.configured_input(),
                             script_directory / "output/custom/summary.csv")

    def test_empty_configuration_has_actionable_error(self):
        with patch.object(converter, "SUMMARY_CSV", " "):
            with self.assertRaisesRegex(ValueError, "SUMMARY_CSV"):
                converter.configured_input()

    def test_import_is_standalone_and_does_not_require_openpyxl(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            script = folder / "summary_to_excel.py"
            shutil.copyfile(converter.__file__, script)
            command = (
                "import builtins, runpy\n"
                "original = builtins.__import__\n"
                "blocked = {'openpyxl', 'pandas', 'numpy', 'matplotlib', 'torch', "
                "'transformers', 'llava', 'llava_pruning', 'run', 'benchmark_report'}\n"
                "def guarded(name, *args, **kwargs):\n"
                "    if name.split('.')[0] in blocked:\n"
                "        raise AssertionError('Unexpected import: ' + name)\n"
                "    return original(name, *args, **kwargs)\n"
                "builtins.__import__ = guarded\n"
                f"runpy.run_path({str(script)!r}, run_name='isolated_import')\n"
            )
            result = subprocess.run([sys.executable, "-c", command], cwd=folder,
                                    capture_output=True, text=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(list(folder.glob("*.xlsx")), [])


@unittest.skipUnless(HAS_OPENPYXL, "requires openpyxl")
class SummaryExcelTests(unittest.TestCase):
    def summary(self, folder, headers=None, rows=None):
        """Write a synthetic CSV fixture; no actual experiment data is used."""
        source = folder / "summary.csv"
        if headers is None:
            headers = ["prune_rate", "complete", "accuracy", "precision", "recall", "tnr"]
        if rows is None:
            rows = [[90, False, 0.875, "", 0.5, 1], [0, True, 1, 1, 1, 1]]
        with source.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
        return source

    def test_preserves_columns_row_order_missing_values_and_incomplete_rows(self):
        from openpyxl import load_workbook

        headers = ["sample_id", "tnr", "prune_rate", "complete", "accuracy", "precision",
                   "recall", "note", "custom_seconds", "custom_number"]
        rows = [["0007", 1, 90, False, 0.875, "", 0.5, "  unchanged  ", "1.25", "001.50"],
                ["0002", 0, 0, True, 0, 1, "", "=SUM(A1:A2)", "", "1e3"]]
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = self.summary(folder, headers, rows)
            original = source.read_bytes()
            destination = converter.convert_summary(source)
            self.assertIsInstance(destination, Path)
            self.assertEqual(destination, (folder / "summary.xlsx").resolve())
            self.assertEqual(source.read_bytes(), original)
            workbook = load_workbook(destination, data_only=False)
            try:
                self.assertEqual(len(workbook.worksheets), 1)
                sheet = workbook.active
                self.assertEqual(sheet.max_column, len(headers))
                self.assertEqual(sheet.max_row, len(rows) + 1)
                self.assertEqual([cell.value for cell in sheet[1]], headers)
                self.assertEqual([cell.value for cell in sheet[2]],
                                 ["0007", 1, 90, False, 0.875, None, 0.5,
                                  "  unchanged  ", 1.25, "001.50"])
                self.assertEqual([cell.value for cell in sheet[3]],
                                 ["0002", 0, 0, True, 0, 1, None,
                                  "=SUM(A1:A2)", None, "1e3"])
                self.assertIs(type(sheet["D2"].value), bool)
                self.assertIs(type(sheet["C2"].value), int)
                for column in ("B", "E", "F", "G"):
                    for row_number in (2, 3):
                        self.assertEqual(sheet[f"{column}{row_number}"].number_format, "0.00%")
                self.assertEqual(sheet["C2"].number_format, '0"%"')
                self.assertEqual(sheet["H3"].data_type, "s")
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:J3")
                self.assertTrue(all(cell.font.name == "Arial"
                                    for row in sheet for cell in row))
                self.assertEqual(sheet._charts, [])
            finally:
                workbook.close()

    def test_known_counts_are_integers_and_unknown_text_is_never_a_formula(self):
        from openpyxl import load_workbook

        counts = ["expected_samples", "evaluated_samples", "labeled_samples", "correct",
                  "incorrect", "unparsed", "parsed_samples", "tp", "fp", "tn", "fn"]
        literal_text = ["=1+1", "+SUM(A1:A2)", "-123", "@SUM(A1:A2)"]
        headers = ["prune_rate", "complete", *counts, "unknown"]
        rows = [[index, True, *range(len(counts)), text]
                for index, text in enumerate(literal_text)]
        with tempfile.TemporaryDirectory() as directory:
            source = self.summary(Path(directory), headers, rows)
            workbook = load_workbook(converter.convert_summary(source), data_only=False)
            try:
                sheet = workbook.active
                for row_number, text in enumerate(literal_text, start=2):
                    for column_number, expected in enumerate(range(len(counts)), start=3):
                        value = sheet.cell(row_number, column_number).value
                        self.assertIs(type(value), int)
                        self.assertEqual(value, expected)
                    cell = sheet.cell(row_number, len(headers))
                    self.assertEqual(cell.value, text)
                    self.assertEqual(cell.data_type, "s")
            finally:
                workbook.close()

    def test_repeated_default_exports_preserve_all_existing_files(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = self.summary(folder)
            (folder / "predictions.jsonl").write_bytes(b"synthetic predictions\n")
            (folder / "metrics.png").write_bytes(b"existing figure")
            before = {path.name: path.read_bytes() for path in folder.iterdir()}
            first = converter.convert_summary(folder)
            first_bytes = first.read_bytes()
            second = converter.convert_summary(source)
            second_bytes = second.read_bytes()
            third = converter.convert_summary(source)
            self.assertEqual([path.name for path in (first, second, third)],
                             ["summary.xlsx", "summary_2.xlsx", "summary_3.xlsx"])
            self.assertEqual(first.read_bytes(), first_bytes)
            self.assertEqual(second.read_bytes(), second_bytes)
            for name, content in before.items():
                self.assertEqual((folder / name).read_bytes(), content)

    def test_explicit_output_never_overwrites_existing_destination_or_source(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = self.summary(folder)
            destination = folder / "chosen.xlsx"
            self.assertEqual(converter.convert_summary(source, destination), destination.resolve())
            before = {path.name: path.read_bytes() for path in folder.iterdir()}
            with self.assertRaises(FileExistsError):
                converter.convert_summary(source, destination)
            with self.assertRaises((ValueError, FileExistsError)):
                converter.convert_summary(source, source)
            self.assertEqual({path.name: path.read_bytes() for path in folder.iterdir()}, before)

    def test_invalid_csv_is_rejected_before_creating_a_workbook(self):
        invalid_csvs = [
            "", "\n", "accuracy\n", "accuracy,accuracy\n0.2,0.3\n",
            "accuracy,\n0.2,text\n", "accuracy,   \n0.2,text\n",
            "accuracy,note\n0.5\n", "accuracy,note\n0.5,text,extra\n",
            "accuracy\n0.5\nnot-a-number\n",
        ]
        invalid_values = {
            "accuracy": ["nan", "inf", "-0.01", "1.01", "90", "abc"],
            "precision": ["NaN", "-1", "2"],
            "recall": ["-inf", "1.1"],
            "tnr": ["inf", "-0.1"],
            "prune_rate": ["-1", "100", "1.5", "nan"],
            "evaluated_samples": ["-1", "1.5", "inf"],
            "tp": ["-1", "1.5"],
            "complete": ["maybe"],
            "decode_seconds": ["nan", "inf", "-inf", "abc"],
        }
        invalid_csvs.extend(f"{column}\n{value}\n"
                            for column, values in invalid_values.items() for value in values)
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = folder / "summary.csv"
            for content in invalid_csvs:
                with self.subTest(content=content):
                    source.write_text(content, encoding="utf-8")
                    original = source.read_bytes()
                    with self.assertRaises(ValueError):
                        converter.convert_summary(source)
                    self.assertEqual(source.read_bytes(), original)
                    self.assertEqual(list(folder.glob("*.xlsx")), [])

    def test_missing_input_does_not_write(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            with self.assertRaises(FileNotFoundError):
                converter.convert_summary(folder)
            self.assertEqual(list(folder.iterdir()), [])

    def test_no_argument_main_uses_top_level_configuration(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = self.summary(folder)
            with patch.object(converter, "SUMMARY_CSV", str(source)):
                with contextlib.redirect_stdout(io.StringIO()):
                    converter.main([])
            self.assertTrue((folder / "summary.xlsx").is_file())

    def test_copied_single_file_cli_accepts_directory_and_explicit_output(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = self.summary(folder)
            original = source.read_bytes()
            script = folder / "summary_to_excel.py"
            shutil.copyfile(converter.__file__, script)
            output = folder / "cli.xlsx"
            command = (
                "import builtins, runpy, sys\n"
                "original = builtins.__import__\n"
                "blocked = {'pandas', 'numpy', 'matplotlib', 'torch', 'transformers', "
                "'llava', 'llava_pruning', 'run', 'benchmark_report'}\n"
                "def guarded(name, *args, **kwargs):\n"
                "    if name.split('.')[0] in {'numpy', 'pandas'}:\n"
                "        raise ImportError('Optional package absent: ' + name)\n"
                "    if name.split('.')[0] in blocked:\n"
                "        raise AssertionError('Unexpected import: ' + name)\n"
                "    return original(name, *args, **kwargs)\n"
                "builtins.__import__ = guarded\n"
                f"sys.argv = [{str(script)!r}, {str(folder)!r}, '--output', {str(output)!r}]\n"
                f"runpy.run_path({str(script)!r}, run_name='__main__')\n"
            )
            result = subprocess.run([sys.executable, "-c", command], cwd=folder,
                                    capture_output=True, text=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(output.is_file())
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
